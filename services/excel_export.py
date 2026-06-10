import io
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter
THIN = Side(style='thin')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEADER_FILL = PatternFill('solid', fgColor='D9E1F2') 
PERIOD_FILL = PatternFill('solid', fgColor='E2EFDA') 
TOTAL_FILL  = PatternFill('solid', fgColor='BDD7EE') 

DEPT_NAMES = {
    'ДМИС': 'Департамент менеджмента и инноваций в спорте',
    'ДСГД': 'Департамент социально-гуманитарных дисциплин',
    'ДСОК': 'Департамент спортивного образования и коучинга',
}

def _cell(ws, row, col, value='', bold=False, align='left', fill=None, border=False):
    c = ws.cell(row=row, column=col, value=value)
    c.font = Font(name='Times New Roman', bold=bold, size=11)
    c.alignment = Alignment(horizontal=align, vertical='center', wrap_text=True)
    if fill: c.fill = fill
    if border: c.border = BORDER
    return c

def generate_teacher_xlsx(teacher, disciplines, director_name='') -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = 'Нагрузка'
    
    col_widths = [3, 5, 45, 15, 12, 8, 10, 10, 8, 10, 10, 10, 8, 10, 8, 8, 10]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.merge_cells('A1:Q1')
    _cell(ws, 1, 1, 'КАЗАХСКИЙ НАЦИОНАЛЬНЫЙ УНИВЕРСИТЕТ СПОРТА', bold=True, align='center')
    
    ws.merge_cells('I2:Q5')
    approve_text = f"УТВЕРЖДАЮ\nДЕКАН\n__________ Ж. Кенжин\n«___» ________ 2025 г."
    c = ws.cell(row=2, column=9, value=approve_text)
    c.alignment = Alignment(horizontal='center', vertical='top', wrap_text=True)

    _cell(ws, 7, 1, 'ИНДИВИДУАЛЬНАЯ ПЕДАГОГИЧЕСКАЯ НАГРУЗКА', bold=True, align='center')
    ws.merge_cells('A7:Q7')

    info = [
        ('Фамилия, имя, отчество', teacher.name),
        ('Должность', teacher.pos),
        ('Ученая степень', teacher.degree),
        ('шт. единицы в кредитах', f"{sum(d.cred for d in disciplines)} кр.")
    ]
    curr_r = 9
    for label, val in info:
        ws.merge_cells(f'A{curr_r}:G{curr_r}')
        _cell(ws, curr_r, 1, label)
        ws.merge_cells(f'H{curr_r}:Q{curr_r}')
        _cell(ws, curr_r, 8, val, bold=True)
        curr_r += 1

    r = curr_r + 2
   
    by_period = {}
    for d in disciplines:
        for p in d.periods: 
            by_period.setdefault(p['num'], []).append((d, p))

    grand_total_hours = 0
    row_idx = 1
    r = 18 

    for p_num in sorted(by_period.keys()):
        ws.merge_cells(f'B{r}:Q{r}')
        _cell(ws, r, 2, f'{p_num} академический период', bold=True, fill=PERIOD_FILL)
        r += 1

        period_h = 0
        for disc, p in by_period[p_num]:
            streams = 1 
            part_cred = float(disc.cred or 0) / 2
            
            vals = [
                '', row_idx, disc.disc, 'СТ-24', 'Русский', disc.cred,
                25, streams, part_cred, part_cred * streams, # Лекции
                25, streams, part_cred, part_cred * streams, # Практики
                p.get('sro', 0), disc.exam_period or '', p['total']
            ]
            
            for ci, v in enumerate(vals, 1):
                _cell(ws, r, ci, v, align='center' if ci != 3 else 'left', border=True)
            
            period_h += p['total']
            row_idx += 1
            r += 1

        ws.merge_cells(f'B{r}:P{r}')
        _cell(ws, r, 2, f'Всего за {p_num} период (час):', bold=True, align='right')
        _cell(ws, r, 17, period_h, bold=True, border=True, fill=PERIOD_FILL)
        grand_total_hours += period_h
        r += 1
    r += 2
    dept_full = DEPT_NAMES.get(teacher.dept, teacher.dept)
    _cell(ws, r, 1, f"{teacher.pos} _____________ {teacher.name}")
    r += 2
    _cell(ws, r, 1, f"Директор {dept_full} _____________ {director_name}")

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()